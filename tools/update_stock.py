#!/usr/bin/env python3
"""
მარაგის განახლება საიტზე.

  python tools/update_stock.py                 # data/მარაგი.xlsx → public/products.json
  python tools/update_stock.py --template      # შექმნის data/მარაგი.xlsx-ს ახლანდელი მარაგით (შაბლონი)
  python tools/update_stock.py ფაილი.xlsx --dry-run   # მხოლოდ შემოწმება, არაფერს წერს

Excel-ის სტრუქტურა (ფურცელი „მარაგი“, პირველი სტრიქონი — სათაურები):
  კოდი | ფერის კოდი | მოდელი | ფერი | კატეგორია | ფასი | ფასდაკლებული ფასი | ფოტო | 27 | 28 | ... | 45
  • კატეგორია: ფეხბურთი / სპორტი / ყოველდღიური / ბავშვები
  • ფასდაკლებული ფასი ცარიელი = ფასდაკლების გარეშე
  • ზომის სვეტში — რამდენი წყვილია (ცარიელი ან 0 = არ არის)
  • ფოტო — ბმული (სურვილისამებრ). თუ public/img/<კოდი>-<ფერი>.jpg უკვე არსებობს, ბმული არ სჭირდება.
"""
import json, os, sys, io, re, urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PUB = ROOT / "public"
DEFAULT = ROOT / "data" / "მარაგი.xlsx"
CATS = {"ფეხბურთი": "football", "სპორტი": "sport", "ყოველდღიური": "casual", "ბავშვები": "kids"}
CATS_R = {v: k for k, v in CATS.items()}
FIXED = ["კოდი", "ფერის კოდი", "მოდელი", "ფერი", "კატეგორია", "ფასი", "ფასდაკლებული ფასი", "ფოტო"]

try:
    import openpyxl
except ImportError:
    sys.exit("საჭიროა openpyxl:  pip install openpyxl pillow numpy scipy")


def size_key(v):
    f = float(v)
    return str(int(f)) if f.is_integer() else str(f)


def make_template(path):
    P = json.loads((PUB / "products.json").read_text(encoding="utf-8"))
    sizes = sorted({s for p in P for s in p["sizes"]}, key=float)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "მარაგი"
    ws.append(FIXED + [float(s) if "." in s else int(s) for s in sizes])
    for p in P:
        code, color = p["id"].split("-", 1)
        ws.append([code, color, p["model"], p["color"], CATS_R.get(p["cat"], p["cat"]), p["price"],
                   p["sale"] if p["sale"] < p["price"] else None, ""] + [p["sizes"].get(s) or None for s in sizes])
    from openpyxl.styles import Font, PatternFill
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="111111")
    ws.freeze_panes = "I2"
    for col, w in zip("ABCDEFGH", (9, 9, 22, 30, 13, 8, 11, 30)):
        ws.column_dimensions[col].width = w
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"შაბლონი შეიქმნა: {path}  ({len(P)} პროდუქტი)")


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["მარაგი"] if "მარაგი" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [h for h in FIXED[:7] if h not in head]
    if missing:
        sys.exit("Excel-ში აკლია სვეტები: " + ", ".join(missing))
    ix = {h: head.index(h) for h in FIXED if h in head}
    size_cols = [(i, size_key(h)) for i, h in enumerate(rows[0]) if i not in ix.values() and h not in (None, "") and re.fullmatch(r"\d+(\.\d+)?", str(h).strip())]
    out, errs = [], []
    for n, r in enumerate(rows[1:], start=2):
        if not r or r[ix["კოდი"]] in (None, ""):
            continue
        code, color = str(r[ix["კოდი"]]).strip(), str(r[ix["ფერის კოდი"]] or "").strip()
        pid = f"{code}-{color}" if color else code
        cat = CATS.get(str(r[ix["კატეგორია"]] or "").strip())
        try:
            price = float(r[ix["ფასი"]])
        except (TypeError, ValueError):
            errs.append(f"სტრიქონი {n} ({pid}): ფასი არ არის"); continue
        sale_raw = r[ix["ფასდაკლებული ფასი"]]
        sale = float(sale_raw) if sale_raw not in (None, "") else price
        if not cat:
            errs.append(f"სტრიქონი {n} ({pid}): კატეგორია უნდა იყოს {', '.join(CATS)}"); continue
        if sale > price:
            errs.append(f"სტრიქონი {n} ({pid}): ფასდაკლებული ფასი ფასზე მეტია"); continue
        sizes = {}
        for i, s in size_cols:
            v = r[i] if i < len(r) else None
            if v not in (None, "") and float(v) > 0:
                sizes[s] = int(float(v))
        sale = round(sale, 2)
        out.append({"id": pid, "model": str(r[ix["მოდელი"]]).strip(), "color": str(r[ix["ფერი"]] or "").strip(), "cat": cat,
                    "price": round(price) if float(price).is_integer() else price, "sale": int(sale) if float(sale).is_integer() else sale,
                    "disc": round((1 - sale / price) * 100) if sale < price else 0, "sizes": sizes, "img": f"img/{pid}.jpg",
                    "_photo": str(r[ix["ფოტო"]] or "").strip() if "ფოტო" in ix else ""})
    ids = [p["id"] for p in out]
    dup = {i for i in ids if ids.count(i) > 1}
    if dup:
        errs.append("გამეორებული კოდები: " + ", ".join(sorted(dup)))
    return out, errs


def fetch_photo(url, dest):
    from PIL import Image
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    data = urllib.request.urlopen(req, timeout=30).read()
    im = Image.open(io.BytesIO(data))
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA"); bg = Image.new("RGB", im.size, "white"); bg.paste(im, mask=im.split()[-1]); im = bg
    im = im.convert("RGB"); im.thumbnail((900, 900))
    s = max(im.size); sq = Image.new("RGB", (s, s), "white"); sq.paste(im, ((s - im.size[0]) // 2, (s - im.size[1]) // 2))
    sq.save(dest, quality=84, optimize=True, progressive=True)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    path = Path(args[0]) if args else DEFAULT
    if "--template" in sys.argv:
        return make_template(path)
    if not path.exists():
        sys.exit(f"ფაილი ვერ მოიძებნა: {path}\nშაბლონისთვის გაუშვი:  python tools/update_stock.py --template")
    P, errs = read(path)
    old = {p["id"]: p for p in json.loads((PUB / "products.json").read_text(encoding="utf-8"))} if (PUB / "products.json").exists() else {}

    no_photo = []
    for p in P:
        dest = PUB / p["img"]
        if not dest.exists():
            if p["_photo"] and "--dry-run" not in sys.argv:
                try:
                    fetch_photo(p["_photo"], dest); print(f"  ფოტო ჩამოიტვირთა: {p['id']}")
                except Exception as e:
                    errs.append(f"{p['id']}: ფოტო ვერ ჩამოიტვირთა ({e})")
            if not dest.exists():
                no_photo.append(p["id"])
    for p in P:
        p.pop("_photo", None)

    live = [p for p in P if p["id"] not in no_photo]
    total = sum(sum(p["sizes"].values()) for p in live)
    new = [p["id"] for p in live if p["id"] not in old]
    gone = [i for i in old if i not in {p["id"] for p in live}]
    changed = [p["id"] for p in live if p["id"] in old and (old[p["id"]]["sizes"] != p["sizes"] or old[p["id"]]["sale"] != p["sale"])]

    print(f"\nწაკითხულია: {len(P)} პროდუქტი · საიტზე გამოჩნდება: {len(live)} · სულ {total} წყვილი")
    if new: print("  ახალი:", ", ".join(new))
    if gone: print("  მოიხსნა საიტიდან:", ", ".join(gone))
    if changed: print(f"  შეიცვალა ფასი/მარაგი: {len(changed)}")
    if no_photo: print("  ⚠ ფოტო არ აქვს (არ გამოჩნდება, სანამ public/img/<კოდი>.jpg არ დაემატება):", ", ".join(no_photo))
    if errs:
        print("\nშეცდომები:"); [print("  ✗", e) for e in errs]
        if "--force" not in sys.argv:
            sys.exit("\nproducts.json არ შეიცვალა. გაასწორე Excel ან გაუშვი --force-ით.")
    if "--dry-run" in sys.argv:
        return print("\n--dry-run: არაფერი ჩაწერილა.")
    (PUB / "products.json").write_text(json.dumps(live, ensure_ascii=False, indent=0), encoding="utf-8")
    print("\n✓ public/products.json განახლდა.")
    try:
        sys.path.insert(0, str(ROOT / "tools"))
        import subprocess
        subprocess.run([sys.executable, str(ROOT / "tools" / "make_cutouts.py")], check=False)
    except Exception as e:
        print("cutouts:", e)
    print("შემდეგი ნაბიჯი: ატვირთე საიტი (deploy) — იხ. README.")


if __name__ == "__main__":
    main()
